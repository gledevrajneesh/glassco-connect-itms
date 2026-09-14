from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Glassco_CONNECT_IT_Support_Desk_Project_Plan.docx"

BLUE = "1467E1"
NAVY = "10294A"
MID = "5B6B82"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2FF"
PALE_GREEN = "E9F6EE"
PALE_AMBER = "FFF4DF"
WHITE = "FFFFFF"
LINE = "D5DCE7"
RED = "B42318"


def set_font(run, size=11, bold=False, color="1F2937", italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, bold=False, color="1F2937", size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, header_fill=LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=NAVY, size=9)
        shade(table.rows[0].cells[i], header_fill)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            set_cell_text(row.cells[i], value)
            if len(table.rows) % 2 == 1:
                shade(row.cells[i], "FAFBFC")
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first, rest = text.split(":", 1)
        r = p.add_run(first + ":")
        set_font(r, bold=True)
        r = p.add_run(rest)
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        r = p.add_run(item)
        set_font(r)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_font(r)


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(label.upper())
    set_font(run, size=9, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run = p2.add_run(text)
    set_font(run, size=10.5, color=NAVY)
    set_table_geometry(table, [9360])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MID)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, before, after, color in (
    ("Heading 1", 16, 16, 8, BLUE),
    ("Heading 2", 13, 12, 6, BLUE),
    ("Heading 3", 12, 8, 4, "1F4D78"),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
for list_name in ("List Bullet", "List Bullet 2", "List Number"):
    style = styles[list_name]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.10

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = hp.add_run("GLASSCO CONNECT  |  IT SUPPORT DESK")
set_font(r, size=9, bold=True, color=NAVY)
footer = section.footer
fp = footer.paragraphs[0]
r = fp.add_run("Discussion Draft  |  Controlled project document")
set_font(r, size=8.5, color=MID)
add_page_number(footer.add_paragraph())

# Memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("PROJECT BRIEF & DELIVERY ROADMAP")
set_font(r, size=10, bold=True, color=BLUE)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Glassco CONNECT - IT Support Desk")
set_font(r, size=24, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("Integrated employee support, service operations and IT asset lifecycle coordination")
set_font(r, size=13, color=MID)

metadata = [
    ("Document status", "Discussion Draft - approval required before development"),
    ("Prepared for", "Glassco Laboratories"),
    ("Platform", "Glassco Communication & Coordination Platform"),
    ("Integration", "Glassco CONNECT ITMS and shared platform masters"),
    ("Prepared on", date.today().strftime("%d %B %Y")),
]
add_table(doc, ["Control", "Detail"], metadata, [2700, 6660], header_fill=PALE_BLUE)
add_callout(doc, "Approval gate", "This document defines the proposed roadmap for discussion. No IT Support Desk development will begin until Glassco confirms the scope, workflow, roles, SLA policy and delivery sequence in writing.", fill=PALE_AMBER, accent="9A6700")

add_heading(doc, "1. Executive summary", 1)
add_body(doc, "Glassco CONNECT IT Support Desk will provide one governed system for employees to request IT assistance and for the IT team to triage, assign, resolve, communicate and report on service work. It will reuse the existing Glassco identity, user, department, location and IT asset records rather than creating duplicate masters.")
add_body(doc, "The recommended solution combines an employee-facing portal, an agent workspace, SLA and escalation controls, notifications, knowledge articles, management dashboards and deep ITMS lifecycle integration. The first delivery will be developed and validated on localhost, followed by controlled Firebase deployment after UAT approval.")

add_heading(doc, "2. Product objectives", 1)
add_bullets(doc, [
    "Give every authorised employee a clear, mobile-friendly way to raise and follow IT requests.",
    "Give support agents one actionable queue with ownership, priority, SLA and complete history.",
    "Connect asset incidents, repairs, replacements, loss and damage directly to ITMS records.",
    "Preserve controlled audit, communication and approval evidence for every material change.",
    "Provide management with drill-down service, asset reliability and workload reporting.",
    "Reduce repeated incidents through reusable knowledge, templates and guided ticket creation.",
])

add_heading(doc, "3. Scope boundary", 1)
add_heading(doc, "3.1 Included in the proposed product", 2)
add_bullets(doc, [
    "Employee ticket creation, tracking, comments, attachments, resolution confirmation and reopening.",
    "Agent queue, assignment, Kanban workflow, ticket workspace, internal notes and activity timeline.",
    "Ticket categories, priorities, support groups, status rules, SLA timers and escalations.",
    "ITMS-linked asset incidents, repairs, replacements, vendor gate passes and lifecycle updates.",
    "Email and in-app notifications with configurable templates and delivery history.",
    "Knowledge articles, response templates, dashboards, exports and audit reporting.",
])
add_heading(doc, "3.2 Outside the first approved release", 2)
add_bullets(doc, [
    "External customer support and sales/service ticketing, which remain separate planned applications.",
    "Paid telephony, SMS, WhatsApp or commercial omnichannel integrations.",
    "Remote device control, endpoint monitoring or automated software deployment.",
    "Public anonymous ticket creation or access outside authorised Glassco identities.",
    "Advanced AI automation until sufficient governed ticket history exists.",
])

add_heading(doc, "4. Users and operating roles", 1)
add_table(doc, ["Role", "Primary responsibility", "Key authority"], [
    ("Employee", "Raise and follow personal requests", "Comment, attach evidence, confirm resolution, reopen within policy"),
    ("Department Manager", "Raise or monitor authorised departmental requests", "Department visibility subject to access policy"),
    ("Support Agent", "Diagnose and resolve assigned work", "Update tickets, add internal notes, request information"),
    ("Support Lead", "Control queues, workload and escalations", "Assign, reprioritise, approve cancellation, manage SLA exceptions"),
    ("IT Asset Manager", "Control asset-related fulfilment", "Initiate repair, replacement, custody and gate-pass processes"),
    ("IT Head", "Management escalation and approval", "Oversee critical incidents, exceptions and service performance"),
    ("Administrator", "Configure the controlled system", "Masters, roles, workflows, SLA and notification configuration"),
    ("Auditor", "Review records and evidence", "Read-only access to tickets, history, reports and exports"),
], [1800, 3600, 3960])

add_heading(doc, "5. Employee experience", 1)
add_numbers(doc, [
    "Sign in using the existing Glassco Workspace SSO and shared platform authority.",
    "Select a request type, describe the need and optionally select one or more allocated assets.",
    "Receive an automatically generated ticket code and acknowledgement.",
    "Follow status, SLA expectation, assigned support contact and conversation timeline.",
    "Respond to questions, upload evidence and review the proposed resolution.",
    "Confirm the resolution, reopen within policy or allow controlled auto-closure.",
])
add_callout(doc, "Employee asset profile", "The employee view will show current ITMS allocations and relevant allocation, repair, replacement, loss, damage and escalation history. It will not create a second asset register.")

add_heading(doc, "6. Service desk workspace", 1)
add_bullets(doc, [
    "Unified queue with My tickets, Team tickets, Unassigned, Overdue, Breached and Waiting views.",
    "Filters for ticket code, employee, department, location, category, asset, priority, SLA and assignee.",
    "Kanban board with governed drag-and-drop status changes and mandatory transition reasons where applicable.",
    "Ticket workspace containing requester, affected assets, conversation, internal notes, tasks, approvals, SLA and immutable timeline.",
    "Quick assignment, response templates, related-ticket links and duplicate/recurring incident indicators.",
])

add_heading(doc, "7. Ticket lifecycle", 1)
add_table(doc, ["Stage", "Purpose", "Exit control"], [
    ("New", "Ticket received and acknowledged", "Required fields present"),
    ("Triaged", "Category, impact, urgency and service owner confirmed", "Priority and support group assigned"),
    ("Assigned", "Named agent accepts accountability", "Owner and target dates recorded"),
    ("In progress", "Diagnosis or fulfilment underway", "Work notes and actions retained"),
    ("Waiting", "Work paused for a valid external dependency", "Reason required; SLA pause follows policy"),
    ("Resolved", "Agent records solution and outcome", "Resolution code and summary required"),
    ("User confirmation", "Employee validates service restoration", "Confirm, reopen or auto-close timer"),
    ("Closed", "Controlled final record", "Closure summary and audit retained"),
], [1500, 4300, 3560])
add_body(doc, "Controlled waiting reasons: Waiting for user, Waiting for vendor, Waiting for approval, Waiting for spare/replacement and Scheduled work. Reopened and Cancelled are exception states with reason and authority requirements.")

add_heading(doc, "8. Classification and automatic numbering", 1)
add_table(doc, ["Ticket class", "Example code", "Typical use"], [
    ("Incident", "INC-2026-00001", "Unplanned interruption, failure or degraded service"),
    ("Service request", "REQ-2026-00001", "Standard fulfilment, information or authorised service"),
    ("Access request", "ACC-2026-00001", "Application, account or permission request"),
    ("Security concern", "SEC-2026-00001", "Suspicious email, malware, access concern or data incident"),
    ("Asset case", "AST-2026-00001", "Repair, replacement, loss, theft, damage or custody concern"),
], [1900, 2100, 5360])
add_body(doc, "Proposed categories include hardware, software, network/internet, email and Google Workspace, printer/peripheral, cybersecurity, asset repair, asset replacement, onboarding/offboarding IT work and general IT assistance. Categories and subcategories will be maintained as controlled masters.")

add_heading(doc, "9. Priority and SLA model - recommended baseline", 1)
add_table(doc, ["Priority", "Business interpretation", "Response target", "Resolution target", "Escalation"], [
    ("Critical", "Business-wide outage, severe security or safety impact", "15 minutes", "4 business hours", "Immediate Support Lead and IT Head"),
    ("High", "Major user/group impact with no practical workaround", "1 business hour", "8 business hours", "Support Lead before breach"),
    ("Medium", "Normal incident or request with manageable impact", "4 business hours", "2 business days", "Queue escalation before breach"),
    ("Low", "Minor issue, information or planned request", "1 business day", "5 business days", "Ageing reminder"),
], [1200, 3100, 1500, 1600, 1960])
add_callout(doc, "Decision required", "These targets are recommended starting values only. Glassco must confirm business hours, holidays, SLA pause conditions, escalation recipients and auto-closure timing before development.", fill=PALE_AMBER, accent="9A6700")

add_heading(doc, "10. ITMS integration", 1)
add_table(doc, ["Support event", "ITMS action", "Control result"], [
    ("Hardware incident", "Link affected allocated asset", "Ticket and asset histories remain connected"),
    ("Repair required", "Create or link maintenance/repair record", "Asset status and repair history update"),
    ("Vendor movement", "Generate controlled outward gate pass", "Custody and vendor movement retained"),
    ("Replacement", "Initiate existing replacement workflow", "Approvals and replaced-asset relationship preserved"),
    ("Loss, theft or damage", "Create asset escalation", "Accountability and evidence retained"),
    ("Closure", "Write verified outcome to linked lifecycle history", "One consistent end-to-end record"),
], [2100, 3300, 3960])
add_body(doc, "Integration principle: the Support Desk may initiate or link an ITMS process, but it must not duplicate ITMS approvals, inventory status, custody, repair, retirement or master data.")

add_heading(doc, "11. Communication and notification controls", 1)
add_bullets(doc, [
    "Ticket-created acknowledgement with ticket code and direct link.",
    "Assignment, reassignment, agent reply and waiting-for-user notifications.",
    "SLA reminder, breach escalation and critical-incident alerts.",
    "Resolution confirmation request, reopen acknowledgement and closure summary.",
    "Configurable recipients, templates, reminder timing and escalation rules.",
    "Delivery history showing recipient, channel, template, timestamp and result.",
])
add_body(doc, "Recommended Phase 1 approach: all substantive replies remain in the portal, while email provides notification and a secure direct link. Email-to-ticket and email reply ingestion should be evaluated only after the controlled portal workflow is stable.")

add_heading(doc, "12. Knowledge and service improvement", 1)
add_bullets(doc, [
    "Searchable troubleshooting and how-to articles visible according to audience.",
    "Suggested articles during ticket creation to reduce avoidable requests.",
    "Agent response templates with controlled ownership and review dates.",
    "Convert a successful resolution into a draft article for review and publication.",
    "Article version history, review status and helpful/not-helpful feedback.",
    "Recurring-ticket analysis by issue, asset, model, location and department.",
])

add_heading(doc, "13. Dashboards and reports", 1)
add_table(doc, ["View", "Measures", "Expected behaviour"], [
    ("Operational dashboard", "Open, unassigned, overdue, breached, critical and waiting tickets", "Every KPI drills into filtered records"),
    ("Workload dashboard", "Tickets by agent, team, status, age and priority", "Supports assignment and capacity decisions"),
    ("Service performance", "First response, resolution, SLA, reopen and satisfaction", "Trend and period comparison"),
    ("Asset reliability", "Recurring failures, repair volume, downtime and replacement", "Links to ITMS asset and model history"),
    ("Management analysis", "Category, department, location, vendor and monthly trends", "Charts plus controlled export"),
], [2000, 3900, 3460])

add_heading(doc, "14. Security, governance and non-functional requirements", 1)
add_bullets(doc, [
    "Google Workspace SSO with an active central Glassco access assignment.",
    "Role-based module, queue, department and action permissions with least privilege.",
    "Immutable event history for assignments, status, priority, SLA, approvals and controlled edits.",
    "Mobile-friendly responsive UI with frozen headers where long operational lists are used.",
    "Global search across tickets, employees, departments, assets, serial numbers and related records.",
    "Controlled export and backup package covering masters, tickets, conversations, relationships and audit trails.",
    "Data-quality checks for orphan references, inactive assignees, invalid asset links and incomplete closures.",
    "Localhost-first development; Firebase deployment only after verified UAT and security-rule review.",
])

add_heading(doc, "15. Proposed delivery roadmap", 1)
add_table(doc, ["Phase", "Delivery focus", "Principal outputs", "Approval gate"], [
    ("1. Foundation", "Masters, roles and controlled workflow", "Ticket codes, categories, priorities, statuses, groups, SLA and audit model", "Foundation design approved"),
    ("2. Employee portal", "Simple self-service experience", "Create ticket, personal list, asset selection, comments, attachments and confirmation", "Employee UAT passed"),
    ("3. Agent workspace", "Queue and accountable execution", "Assignment, filters, Kanban, detail workspace, internal notes and timeline", "Service desk UAT passed"),
    ("4. ITMS integration", "One asset-service lifecycle", "Repair, replacement, gate pass, escalation and lifecycle updates", "ITMS controls verified"),
    ("5. SLA & communication", "Time and communication governance", "Timers, reminders, escalations, templates and delivery history", "SLA policy approved"),
    ("6. Knowledge & automation", "Reuse and assisted resolution", "Articles, templates, suggestions and recurring-issue detection", "Knowledge governance approved"),
    ("7. Reporting & go-live", "Management visibility and release", "Dashboards, exports, audit, data quality, UAT closure and Firebase deployment", "Go-live approval"),
], [1200, 2200, 3900, 2060])

add_heading(doc, "16. Indicative acceptance criteria", 1)
add_bullets(doc, [
    "An authorised employee can create a valid ticket and receive an automatic code and acknowledgement.",
    "An agent can triage, assign, work, communicate, resolve and close a ticket with a complete timeline.",
    "Invalid status changes, missing mandatory reasons and unauthorised actions are blocked.",
    "SLA targets, pauses, reminders, breaches and escalations calculate according to approved policy.",
    "A linked asset ticket displays accurate ITMS ownership and can initiate approved lifecycle actions without duplication.",
    "Every dashboard KPI opens its underlying filtered records.",
    "Notification delivery and access changes are traceable.",
    "Desktop and mobile UAT passes with no overlap, clipping or inaccessible controls.",
    "Controlled export and restoration are tested before production reliance.",
])

add_heading(doc, "17. Dependencies and risks", 1)
add_table(doc, ["Item", "Risk", "Planned control"], [
    ("Shared masters", "Duplicate or inconsistent employee/department data", "Reuse Glassco Connect shared masters only"),
    ("ITMS integration", "Ticket actions could conflict with asset lifecycle", "Use ITMS as source of truth and call existing workflows"),
    ("SLA design", "Unrealistic targets create false breaches", "Approve calendar, targets and pause rules before build"),
    ("Email delivery", "Free-plan limitations or unreliable delivery", "Portal-first communications and tested notification history"),
    ("Attachments", "Evidence may remain device-local", "Select approved Firebase storage design before live reliance"),
    ("Permissions", "Excess visibility or unauthorised status changes", "Role, queue and department rules tested with real personas"),
    ("Data migration", "Existing support history may be incomplete", "Use validated templates and staged import with reconciliation"),
], [1900, 3300, 4160])

add_heading(doc, "18. Decisions required before development", 1)
decision_rows = [
    ("Application architecture", "Second application on Glassco Connect sharing identity and masters", "Recommended"),
    ("Who may raise for others", "Employee for self; manager for authorised department", "Confirm"),
    ("Employee closure", "Confirmation required; controlled auto-close after waiting period", "Confirm period"),
    ("Ticket cancellation", "Support Lead or Administrator only", "Confirm"),
    ("Critical escalation", "Immediate Support Lead and IT Head notification", "Confirm"),
    ("SLA calendar", "Glassco business hours, holidays and pause reasons", "Define"),
    ("Attachments", "Optional generally; required for selected damage/theft cases", "Confirm"),
    ("Email replies", "Portal replies in Phase 1; email ingestion later", "Confirm"),
    ("Phase sequence", "Seven phases shown in this document", "Approve or amend"),
]
add_table(doc, ["Decision", "Recommended position", "Status"], decision_rows, [2400, 5260, 1700], header_fill=PALE_AMBER)

add_heading(doc, "19. Approval and next step", 1)
add_callout(doc, "Development hold", "The next activity is a structured roadmap review with Glassco. Development begins only after the authorised project sponsor confirms the required decisions and explicitly instructs the team to proceed.", fill=PALE_GREEN, accent="087A3B")
add_table(doc, ["Approval role", "Name", "Decision", "Date"], [
    ("Project sponsor", "", "Approve / Amend / Reject", ""),
    ("IT Head", "", "Approve / Amend / Reject", ""),
    ("IT Asset / Support Manager", "", "Approve / Amend / Reject", ""),
], [2200, 2600, 2760, 1800])

add_heading(doc, "Appendix A - Recommended discussion sequence", 1)
add_numbers(doc, [
    "Confirm application architecture, shared masters and user groups.",
    "Confirm ticket classes, categories, statuses and automatic numbering.",
    "Confirm roles, visibility boundaries and approval authority.",
    "Confirm priority, SLA, business hours, pause and escalation rules.",
    "Confirm ITMS integration events and source-of-truth boundaries.",
    "Confirm notification, attachment and evidence requirements.",
    "Confirm roadmap sequence, UAT participants and go-live gate.",
])

doc.core_properties.title = "Glassco CONNECT IT Support Desk - Project Brief and Delivery Roadmap"
doc.core_properties.subject = "Discussion draft for scope and roadmap approval"
doc.core_properties.author = "Glassco Laboratories"
doc.core_properties.keywords = "Glassco, IT Support Desk, roadmap, ITMS, service management"
doc.core_properties.comments = "Prepared for discussion; development requires explicit approval."
doc.save(OUTPUT)
print(OUTPUT)
